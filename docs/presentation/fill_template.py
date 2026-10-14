"""Fill the Sri Eshwar template (slides 1-8) with the EOPS paper content."""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1F, 0x4E, 0x79)
DEEP = RGBColor(0x14, 0x36, 0x55)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
BORDER = RGBColor(0xD3, 0xDD, 0xE6)
AMBER = RGBColor(0xB5, 0x6A, 0x00)
GREEN = RGBColor(0x2A, 0x6F, 0x52)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation("template.pptx")
slides = prs.slides


def drop_slides_after(n):
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides)[n:]:
        rId = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        prs.part.drop_rel(rId)
        xml_slides.remove(sld)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def shape_named(slide, needle):
    for sh in slide.shapes:
        if sh.has_text_frame and needle.lower() in sh.text_frame.text.lower():
            return sh
    return None


def keep_heading_only(slide):
    """Strip the template's authoring instructions, keep only the section heading."""
    for sh in list(slide.shapes):
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        text = tf.text.strip()
        if text.lower().startswith("slide no."):
            remove(sh)
            continue
        if not text:
            continue
        paras = list(tf.paragraphs)
        keep = next((p for p in paras if p.text.strip()), None)
        if keep is None or len(paras) == 1:
            continue
        for para in paras:  # never delete the one paragraph we keep
            if para is not keep:
                para._p.getparent().remove(para._p)
        if sh.width and sh.height:
            sh.left, sh.width = sh.left, sh.width
            sh.top, sh.height = Inches(0.3), Inches(0.9)
    for sh in list(slide.shapes):  # empty content placeholder: we draw our own
        if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
            remove(sh)


def textbox(slide, x, y, w, h, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
            italic=False, anchor=MSO_ANCHOR.TOP, space_after=6, line=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    box._defaults = dict(size=size, bold=bold, color=color, align=align, italic=italic,
                         space_after=space_after, line=line)
    return box


def write(box, items):
    """items: str | (text, overrides) | list of runs [(text, overrides)]"""
    d = box._defaults
    tf = box.text_frame
    first = True
    for item in items:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        runs = item if isinstance(item, list) else [item]
        over = {}
        for run_spec in runs:
            text, o = (run_spec, {}) if isinstance(run_spec, str) else run_spec
            over = o
            r = para.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", d["size"]))
            f.bold = o.get("bold", d["bold"])
            f.italic = o.get("italic", d["italic"])
            f.color.rgb = o.get("color", d["color"])
        para.alignment = over.get("align", d["align"])
        para.space_after = Pt(over.get("space_after", d["space_after"]))
        if d["line"]:
            para.line_spacing = d["line"]
    return box


def card(slide, x, y, w, h, fill=CARD, border=BORDER, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    if radius:
        shape.adjustments[0] = 0.08
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    return shape


def label(shape, text, *, size=12, bold=True, color=INK, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = align
    for i, line in enumerate(text.split("\n")):
        p = para if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        r.font.color.rgb = color


def arrow(slide, x, y, w, h=0.0):
    conn = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h or 0.14))
    conn.fill.solid()
    conn.fill.fore_color.rgb = RGBColor(0x9A, 0xA8, 0xB4)
    conn.line.fill.background()
    conn.shadow.inherit = False
    return conn


def down_arrow(slide, x, y, w, h):
    conn = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    conn.fill.solid()
    conn.fill.fore_color.rgb = RGBColor(0x9A, 0xA8, 0xB4)
    conn.line.fill.background()
    conn.shadow.inherit = False
    return conn


def bullets(slide, x, y, w, h, items, *, size=14, space=8, color=INK):
    box = textbox(slide, x, y, w, h, size=size, color=color, space_after=space)
    write(box, [[("•  ", {"color": NAVY, "bold": True})] + (it if isinstance(it, list) else [it])
                for it in items])
    return box


def set_runs(shape, text, *, size=None, bold=None, color=None):
    """Replace a template text frame's content, keeping its formatting."""
    tf = shape.text_frame
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    para = tf.paragraphs[0]
    for r in list(para.runs)[1:]:
        r._r.getparent().remove(r._r)
    if not para.runs:
        para.add_run()
    run = para.runs[0]
    run.text = text
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return run


# ───────────────────────── slide 1 — title ─────────────────────────
s1 = slides[0]
title = s1.shapes.title
set_runs(title, "EOPS: A Multi-Agent AI Engineering Operations Platform for "
                "Automated Incident Investigation and Root Cause Analysis", size=24, bold=True, color=DEEP)
title.left, title.top = Inches(0.6), Inches(1.6)
title.width, title.height = Inches(8.8), Inches(2.1)

for sh in s1.shapes:
    if sh.has_text_frame and "Students Name" in sh.text_frame.text:
        tf = sh.text_frame
        for p in list(tf.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        p = tf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        for i, line in enumerate([("Sanjay P C", True), ("Sam Benadick J", False),
                                  ("Praveenkumar S", False), ("Nibuna A", False),
                                  ("B.Tech AI & Data Science", False)]):
            para = p if i == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            r = para.add_run()
            r.text = line[0]
            r.font.name, r.font.size, r.font.bold = FONT, Pt(15 if i < 4 else 12), line[1]
            r.font.color.rgb = INK if i < 4 else GREY
    elif sh.has_text_frame and "Guide Name" in sh.text_frame.text:
        tf = sh.text_frame
        for p in list(tf.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        p = tf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        for i, (line, sz, bold) in enumerate([("Ms D. Hemamalini", 15, True),
                                              ("Assistant Professor", 12, False),
                                              ("Dept. of AI & Data Science", 12, False),
                                              ("Sri Eshwar College of Engineering", 12, False)]):
            para = p if i == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            r = para.add_run()
            r.text = line
            r.font.name, r.font.size, r.font.bold = FONT, Pt(sz), bold
            r.font.color.rgb = INK if i == 0 else GREY
    elif sh.is_placeholder and sh.placeholder_format.type == 15:  # footer
        set_runs(sh, "ICSIE 2026 — Paper Presentation")

# ───────────────────────── slide 2 — introduction ─────────────────────────
s = slides[1]
keep_heading_only(s)

stat = card(s, 6.85, 1.6, 2.65, 1.9, fill=DEEP, border=DEEP)
num = textbox(s, 6.95, 1.78, 2.45, 0.6, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
write(num, ["52.5 h"])
cap = textbox(s, 6.95, 2.48, 2.45, 0.9, size=10.5, color=RGBColor(0xD6, 0xE2, 0xEA), align=PP_ALIGN.CENTER)
write(cap, ["median time to mitigate a production incident (Microsoft, one year of incidents)"])

bullets(s, 0.55, 1.62, 6.1, 3.0, [
    [("Cloud-native systems ship daily; ", {}), ("production incidents are unavoidable.", {})],
    [("One investigation spans six disconnected systems: ", {}),
     ("source code, deployments, database, logs, runbooks and team chat.", {"bold": True})],
    "Engineers correlate this evidence manually. Context switching raises cognitive load, prolongs Mean Time to Resolution and makes diagnosis depend on individual expertise.",
    "Existing LLM assistants target code generation or conversation, not coordinated investigation across engineering tools.",
], size=13.5, space=12)

gap = card(s, 0.55, 4.75, 8.95, 1.2, fill=RGBColor(0xFB, 0xF2, 0xE2), border=RGBColor(0xEB, 0xD9, 0xBC))
obj = textbox(s, 0.8, 4.95, 8.5, 0.95, size=13)
write(obj, [[("Objective:  ", {"bold": True, "color": AMBER}),
             ("a supervisor-driven multi-agent platform where specialised agents investigate a team's own tools in parallel and return an evidence-cited root cause, with human approval before any action.", {})]])

# ───────────────────────── slide 3 — literature review ─────────────────────────
s = slides[2]
keep_heading_only(s)

rows = [
    ("Work", "Contribution", "Limitation for incident investigation"),
    ("Lewis et al., NeurIPS 2020 — RAG", "Grounds generation in retrieved documents", "Retrieval only; no tools, no investigation"),
    ("Yao et al., ICLR 2023 — ReAct", "Interleaves reasoning with tool actions", "Single agent on a single task"),
    ("Wu et al., COLM 2024 — AutoGen", "Multi-agent conversation framework", "Unbounded chat; no evidence contract"),
    ("Ahmed et al., ICSE 2023", "Fine-tuned LLM proposes root cause from incident text", "One-shot; no multi-source evidence"),
    ("Chen et al., EuroSys 2024 — RCACopilot", "Handlers collect diagnostics; LLM predicts category (Micro-F1 0.766)", "Provider-side; one handler per alert type"),
    ("Roy et al., FSE 2024", "ReAct agent for RCA; hallucination 6% vs 49% for retrieval", "Single agent; diagnosis only"),
    ("Yang et al., NeurIPS 2024 — SWE-agent", "Agent-computer interface resolves repository issues", "Repository only; no live system signals"),
]
table_shape = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.5), Inches(9.0), Inches(3.5))
tbl = table_shape.table
tbl.columns[0].width = Inches(2.75)
tbl.columns[1].width = Inches(3.55)
tbl.columns[2].width = Inches(2.70)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(0.42)
    for c, text in enumerate(row):
        cell = tbl.cell(r, c)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = DEEP if r == 0 else (WHITE if r % 2 else CARD)
        para = cell.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(10.5 if r else 11)
        run.font.bold = r == 0
        run.font.color.rgb = WHITE if r == 0 else INK

gapc = card(s, 0.5, 5.25, 9.0, 1.2, fill=DEEP, border=DEEP)
gt = textbox(s, 0.75, 5.45, 8.5, 0.9, size=12.5, color=WHITE)
write(gt, [[("Research gap:  ", {"bold": True, "color": RGBColor(0xF0, 0xB4, 0x5F)}),
            ("no system coordinates several specialised agents over a customer's own heterogeneous tools, "
             "sharing one investigation state in which every finding must cite the tool call that produced it.", {})]])

# ───────────────────────── slide 4 — methodology ─────────────────────────
s = slides[3]
keep_heading_only(s)

inc = card(s, 0.5, 1.45, 1.9, 0.55, fill=CARD)
label(inc, "Incident in chat", size=11)
arrow(s, 2.5, 1.63, 0.35, 0.18)
sup = card(s, 2.95, 1.45, 2.2, 0.55, fill=NAVY, border=NAVY)
label(sup, "Supervisor plans", size=11, color=WHITE)
note = textbox(s, 5.35, 1.45, 4.15, 0.6, size=10.5, color=GREY, italic=True, anchor=MSO_ANCHOR.MIDDLE)
write(note, ["may assign only desks that have a connected tool"])
down_arrow(s, 3.95, 2.05, 0.2, 0.3)

desks = ["Code", "Frontend\n& Hosting", "Backend\n& Services", "Database", "Observability", "Knowledge"]
for i, d in enumerate(desks):
    c = card(s, 0.5 + i * 1.53, 2.4, 1.4, 0.72, fill=RGBColor(0xE4, 0xED, 0xF4))
    label(c, d, size=9.5, color=DEEP)
run_note = textbox(s, 0.5, 3.17, 9.0, 0.25, size=10, color=GREY, italic=True, align=PP_ALIGN.CENTER)
write(run_note, ["six desks run concurrently · each holds read-only tools · a desk may ask one colleague a precise question"])

state = card(s, 0.5, 3.5, 9.0, 0.5, fill=RGBColor(0xFB, 0xF2, 0xE2), border=RGBColor(0xEB, 0xD9, 0xBC))
label(state, "Shared investigation state — every finding carries its tool-call id", size=11.5)
down_arrow(s, 4.85, 4.02, 0.2, 0.26)

rca = card(s, 1.55, 4.33, 3.3, 0.6, fill=RGBColor(0x1C, 0x72, 0x93), border=RGBColor(0x1C, 0x72, 0x93))
label(rca, "RCA agent — cause, failure point, confidence", size=10.5, color=WHITE)
arrow(s, 4.95, 4.55, 0.3, 0.16)
hum = card(s, 5.35, 4.33, 3.1, 0.6, fill=GREEN, border=GREEN)
label(hum, "Engineer approves before any write", size=10.5, color=WHITE)

phases = ["Initialise", "Plan", "Parallel evidence", "Collaborative reasoning", "Human validation"]
for i, p in enumerate(phases):
    x = 0.5 + i * 1.82
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(5.2), Inches(0.3), Inches(0.3))
    dot.fill.solid(); dot.fill.fore_color.rgb = NAVY; dot.line.fill.background(); dot.shadow.inherit = False
    label(dot, str(i + 1), size=10, color=WHITE)
    t = textbox(s, x + 0.36, 5.22, 1.4, 0.55, size=10, color=INK)
    write(t, [p])
ph = textbox(s, 0.5, 5.75, 9.0, 0.7, size=11, color=GREY)
write(ph, ["Five phases of the proposed methodology; each desk runs a budgeted reason-act loop "
           "(8 tool calls, 300 s) and is forced to report its findings when a budget is reached."])

# ───────────────────────── slide 5 — experimental analysis ─────────────────────────
s = slides[4]
keep_heading_only(s)

exps = [
    ("Experiment 1 — Orchestration", "Does the supervisor coordinate desks and keep one consistent state?",
     "Orchestration latency · state consistency · agent completion rate"),
    ("Experiment 2 — Knowledge retrieval", "How well does the Engineering Memory Layer retrieve runbooks and postmortems?",
     "MRR · Recall@K · Precision@K · response latency"),
    ("Experiment 3 — Root-cause reasoning", "Do generated reports match expert analyses of the same incident?",
     "Evidence use · completeness · explanation quality · confidence calibration"),
    ("Experiment 4 — End-to-end workflow", "Injected faults: API latency, slow query, failed deploy, broken content.",
     "Time to root cause · evidence completeness · unsupported-claim rate · usability"),
]
for i, (head, q, metrics) in enumerate(exps):
    col, row = i % 2, i // 2
    x, y = 0.5 + col * 4.6, 1.5 + row * 1.72
    card(s, x, y, 4.4, 1.55)
    h = textbox(s, x + 0.2, y + 0.14, 4.0, 0.3, size=12, bold=True, color=NAVY)
    write(h, [head])
    b = textbox(s, x + 0.2, y + 0.46, 4.0, 0.55, size=10.5, color=INK)
    write(b, [q])
    m = textbox(s, x + 0.2, y + 1.02, 4.0, 0.45, size=10, color=GREY, italic=True)
    write(m, [metrics])

base = card(s, 0.5, 4.98, 9.0, 1.5, fill=DEEP, border=DEEP)
bt = textbox(s, 0.78, 5.15, 8.5, 1.2, size=11.5, color=WHITE, space_after=5)
write(bt, [
    [("Baselines:  ", {"bold": True, "color": RGBColor(0xF0, 0xB4, 0x5F)}),
     ("single ReAct agent with the same tools · one-shot prompt with all context · the same architecture without compaction", {})],
    [("Datasets:  ", {"bold": True, "color": RGBColor(0xF0, 0xB4, 0x5F)}),
     ("OpenRCA (ICLR 2025, 335 real failures) · AIOpsLab fault scenarios (MLSys 2025) · 20–30 faults injected into a live stack", {})],
    [("Ablations:  ", {"bold": True, "color": RGBColor(0xF0, 0xB4, 0x5F)}),
     ("multi-agent vs single agent · with and without retrieval · with and without enforced citation", {})],
])

# ───────────────────────── slide 6 — results & discussion ─────────────────────────
s = slides[5]
keep_heading_only(s)

chart_data = CategoryChartData()
chart_data.categories = ["Without\ncompaction", "With\ncompaction"]
chart_data.add_series("Input characters resent", (288000, 175500))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.45),
                        Inches(4.3), Inches(3.05), chart_data)
chart = gf.chart
chart.has_legend = False
chart.has_title = True
chart.chart_title.text_frame.text = "Context resent per desk (8 tool calls)"
tr = chart.chart_title.text_frame.paragraphs[0].runs[0]
tr.font.size, tr.font.bold, tr.font.name, tr.font.color.rgb = Pt(11), True, FONT, INK
plot = chart.plots[0]
plot.has_data_labels = True
plot.data_labels.number_format = "#,##0"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(10)
plot.data_labels.font.name = FONT
plot.data_labels.font.color.rgb = INK
plot.vary_by_categories = False
plot.series[0].format.fill.solid()
plot.series[0].format.fill.fore_color.rgb = NAVY
for axis in (chart.category_axis, chart.value_axis):
    axis.tick_labels.font.size = Pt(9)
    axis.tick_labels.font.name = FONT
    axis.tick_labels.font.color.rgb = GREY
chart.value_axis.has_major_gridlines = True
chart.value_axis.maximum_scale = 320000
chart.value_axis.minimum_scale = 0
chart.value_axis.major_unit = 80000

findings = [
    ("39%", "less input resent per desk — older tool output collapses to its opening while keeping its citation id"),
    ("Bounded", "context per step: each desk sees only its own tools; the RCA agent reads one-line evidence, never raw logs"),
    ("Parallel", "wall-clock follows the slowest desk rather than the sum of all desks"),
]
for i, (big, text) in enumerate(findings):
    y = 1.5 + i * 1.03
    c = card(s, 5.05, y, 4.45, 0.92)
    bx = textbox(s, 5.22, y + 0.12, 1.25, 0.35, size=15, bold=True, color=NAVY)
    write(bx, [big])
    tx = textbox(s, 5.22, y + 0.44, 4.1, 0.45, size=10, color=INK)
    write(tx, [text])

impl = card(s, 0.5, 4.62, 9.0, 0.72, fill=CARD)
label(impl, "Implemented: 9 agent roles · 3 live connectors (GitHub, Supabase, knowledge vault) · "
            "229 automated tests · pilot investigation ≈ 7.7 min end to end", size=11)

disc = textbox(s, 0.5, 5.5, 9.0, 1.0, size=11.5, color=INK, space_after=5)
write(disc, [
    [("Discussion:  ", {"bold": True, "color": NAVY}),
     ("the architecture behaves as designed — desks work concurrently, every finding carries a verifiable citation, and compaction cuts resent context measurably.", {})],
    [("Threat to validity:  ", {"bold": True, "color": RGBColor(0xB3, 0x26, 0x1E)}),
     ("this is a position paper with a working prototype; the pilot timing is an engineering observation, not an evaluation result. Large-scale measurement is future work.", {})],
])

# ───────────────────────── slide 7 — conclusion ─────────────────────────
s = slides[6]
keep_heading_only(s)

points = [
    ("A coordination protocol, not a bigger model",
     "A supervisor plans, specialised desks investigate the team's own tools in parallel, and one shared state carries every finding with its tool-call citation."),
    ("Trust comes from the harness",
     "Uncited findings are rejected, investigating desks are read-only, budgets bound every loop, and confidence is capped by code — not by the model's own claim."),
    ("Decision support is the right target today",
     "Published agents resolve only 11–15% of real scenarios autonomously, so EOPS keeps the engineer in the approval seat."),
    ("Future work",
     "Measurement on OpenRCA and injected faults, graph-augmented retrieval over code–service–document relations, and learning from accepted versus rejected recommendations."),
]
for i, (head, body) in enumerate(points):
    y = 1.5 + i * 1.22
    card(s, 0.5, y, 9.0, 1.1)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), Inches(y + 0.3), Inches(0.36), Inches(0.36))
    dot.fill.solid(); dot.fill.fore_color.rgb = NAVY; dot.line.fill.background(); dot.shadow.inherit = False
    label(dot, str(i + 1), size=11, color=WHITE)
    h = textbox(s, 1.27, y + 0.16, 8.0, 0.32, size=13, bold=True, color=NAVY)
    write(h, [head])
    b = textbox(s, 1.27, y + 0.5, 8.0, 0.55, size=11, color=INK)
    write(b, [body])

# ───────────────────────── slide 8 — references ─────────────────────────
s = slides[7]
keep_heading_only(s)

refs_left = [
    "[1] A. Vaswani et al., \"Attention is all you need,\" in Proc. NeurIPS, 2017.",
    "[2] P. Lewis et al., \"Retrieval-augmented generation for knowledge-intensive NLP tasks,\" in Proc. NeurIPS, 2020.",
    "[3] S. Yao et al., \"ReAct: Synergizing reasoning and acting in language models,\" in Proc. ICLR, 2023.",
    "[4] Q. Wu et al., \"AutoGen: Enabling next-gen LLM applications via multi-agent conversation,\" in Proc. COLM, 2024.",
    "[5] J. Yang et al., \"SWE-agent: Agent-computer interfaces enable automated software engineering,\" in Proc. NeurIPS, 2024.",
    "[6] C. E. Jimenez et al., \"SWE-bench: Can language models resolve real-world GitHub issues?,\" in Proc. ICLR, 2024.",
    "[7] T. Ahmed et al., \"Recommending root-cause and mitigation steps for cloud incidents using large language models,\" in Proc. ICSE, 2023.",
]
refs_right = [
    "[8] Y. Chen et al., \"Automatic root cause analysis via large language models for cloud incidents,\" in Proc. EuroSys, 2024.",
    "[9] D. Roy et al., \"Exploring LLM-based agents for root cause analysis,\" in Proc. FSE (Industry), 2024.",
    "[10] \"OpenRCA: Can large language models locate the root cause of software failures?,\" in Proc. ICLR, 2025.",
    "[11] \"AIOpsLab: A holistic framework to evaluate AI agents for enabling autonomous clouds,\" in Proc. MLSys, 2025.",
    "[12] \"ITBench: Evaluating AI agents across diverse real-world IT automation tasks,\" in Proc. ICML, 2025.",
    "[13] Y. Gao et al., \"Retrieval-augmented generation for large language models: A survey,\" arXiv:2312.10997, 2023.",
    "[14] LangChain AI, \"LangGraph,\" 2024; Chroma, \"Chroma embedding database,\" 2024.",
]
for col, refs in enumerate((refs_left, refs_right)):
    box = textbox(s, 0.5 + col * 4.62, 1.5, 4.4, 4.7, size=9.5, color=INK, space_after=9)
    write(box, refs)

note = textbox(s, 0.5, 6.1, 9.0, 0.45, size=10, color=GREY, italic=True)
write(note, ["References [8]–[12] situate EOPS against the current state of the art in LLM-based root cause analysis."])

drop_slides_after(8)
prs.save("EOPS_ICSIE_template_filled.pptx")
print("saved; slides =", len(prs.slides._sldIdLst))
